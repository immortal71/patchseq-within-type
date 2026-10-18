"""Assemble the manuscript, build Word and PDF in one text colour, clean metadata.

Usage: python build.py [black|blue]
"""
import re, subprocess, sys, os, zipfile, shutil, datetime
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
os.chdir(HERE)
mode = sys.argv[1] if len(sys.argv) > 1 else "black"
COLOR = {"black": RGBColor(0, 0, 0), "blue": RGBColor(0x1F, 0x3A, 0x93)}[mode]
HEX = str(COLOR)
OUT = "ion_channel_full_article" + ("" if mode == "black" else "_blue")

# 1. assemble
parts = ["front", "intro", "results", "table1", "discussion", "methods", "back"]
text = "\n\n".join(open(p + ".md", encoding="utf-8").read().strip() for p in parts) + "\n"
open("full_article.md", "w", encoding="utf-8").write(text)
figs = {1: "fig1_raw_vs_within", 2: "fig2_calibration", 3: "fig3_hh_parameters",
        4: "fig4_partitions", 5: "fig5_hcn"}
for n, f in figs.items():
    text = text.replace(f"**Figure {n}.", f"![](../figures/{f}.png){{width=100%}}\n\n**Figure {n}.", 1)
open("build.md", "w", encoding="utf-8").write(text)

# 2. hidden characters and dashes check (fails the build if any slipped in)
bad = re.findall(r"[\u2012-\u2015\u200b-\u200f\u202a-\u202e\u2060\ufeff\u00a0]", text)
assert not bad, f"hidden or dash characters found: {set(bad)}"

# 3. reference docx: every style one colour, one serif font
ref = "reference.docx"
with open(ref, "wb") as fh:
    fh.write(subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"],
                            capture_output=True, check=True).stdout)
d = Document(ref)
for st in d.styles:
    try:
        f = st.font
    except AttributeError:
        continue
    if f is None:
        continue
    f.color.rgb = COLOR
    f.name = "Times New Roman"
for name, size in [("Normal", 11), ("Body Text", 11), ("First Paragraph", 11), ("Title", 17),
                   ("Heading 1", 14), ("Heading 2", 13), ("Heading 3", 11.5), ("Caption", 9.5),
                   ("Image Caption", 9.5), ("Table Caption", 9.5), ("Compact", 9)]:
    try:
        d.styles[name].font.size = Pt(size)
    except KeyError:
        pass
for name in ["Heading 1", "Heading 2", "Heading 3", "Title"]:
    try:
        d.styles[name].font.bold = True
        d.styles[name].font.italic = False
    except KeyError:
        pass
d.save(ref)
# theme colours used by hyperlinks and some headings
tmp = ref + ".tmp"
with zipfile.ZipFile(ref) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename.endswith(".xml"):
            s = data.decode("utf-8")
            s = re.sub(r'w:color w:val="[0-9A-Fa-f]{6}"', f'w:color w:val="{HEX}"', s)
            s = re.sub(r'w:themeColor="[a-zA-Z0-9]+"', "", s)
            data = s.encode("utf-8")
        zout.writestr(item, data)
shutil.move(tmp, ref)

# 4. pandoc
subprocess.run(["pandoc", "build.md", "-o", OUT + ".docx", "--reference-doc", ref,
                "--resource-path", ".:../figures"], check=True)

# 5. post-process: table widths, colours on every run, metadata
d = Document(OUT + ".docx")
for t in d.tables:
    widths = [Cm(4.6)] + [Cm(1.45)] * 6 + [Cm(3.2)]
    for row in t.rows:
        for i, c in enumerate(row.cells):
            if i < len(widths):
                c.width = widths[i]
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)
                    r.font.color.rgb = COLOR
for p in d.paragraphs:
    for r in p.runs:
        r.font.color.rgb = COLOR
    if p.style.name == "Title":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
for s in d.sections:
    s.left_margin = s.right_margin = Cm(2.3)
    s.top_margin = s.bottom_margin = Cm(2.2)
cp = d.core_properties
cp.author = "Aashish Kharel"
cp.last_modified_by = "Aashish Kharel"
cp.title = open("front.md", encoding="utf-8").readline().lstrip("# ").strip()
cp.subject = cp.keywords = cp.comments = cp.category = ""
cp.created = cp.modified = datetime.datetime(2026, 10, 17)
d.save(OUT + ".docx")

# hyperlinks in the body: force colour in document.xml too
tmp = OUT + ".tmp"
with zipfile.ZipFile(OUT + ".docx") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in ("word/document.xml", "word/styles.xml"):
            s = data.decode("utf-8")
            s = re.sub(r'w:color w:val="[0-9A-Fa-f]{6}"', f'w:color w:val="{HEX}"', s)
            s = re.sub(r'w:themeColor="[a-zA-Z0-9]+"', "", s)
            s = re.sub(r'<w:rFonts [^>]*/>', '<w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:eastAsia="Times New Roman" w:cs="Times New Roman"/>', s)
            data = s.encode("utf-8")
        if item.filename == "docProps/app.xml":
            data = re.sub(rb"<Application>.*?</Application>", b"<Application>Microsoft Office Word</Application>", data)
        zout.writestr(item, data)
shutil.move(tmp, OUT + ".docx")

# 6. PDF
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", OUT + ".docx"], check=True,
               capture_output=True)
print("built", OUT + ".docx", OUT + ".pdf")
